"""
Exporte le modèle en étoile (schéma bi) vers un classeur Excel unique.
Chaque feuille est une table Excel nommée : Power BI Service la reconnaît
comme table à importer.

Usage (depuis la racine du projet) :
    python -m src.export_powerbi
"""
import datetime as dt
from pathlib import Path

import pandas as pd
from openpyxl import Workbook
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.table import Table, TableStyleInfo

from src.db import get_engine

OUT_PATH = Path("powerbi/data/olist_model.xlsx")

QUERIES = {
    "FactSales": """
        SELECT order_id, order_item_id, order_date, date_key, customer_key, product_key,
               geography_key, seller_id, quantity, revenue, freight_value,
               is_first_purchase_day::int AS is_first_purchase_day
        FROM bi.fact_sales
        ORDER BY order_id, order_item_id""",
    "DimCustomer": "SELECT * FROM bi.dim_customer ORDER BY customer_key",
    "DimProduct": "SELECT * FROM bi.dim_product ORDER BY product_key",
    "DimGeography": "SELECT * FROM bi.dim_geography ORDER BY geography_key",
    "DimDate": 'SELECT * FROM bi.dim_date ORDER BY "Date"',
}

# Valeurs validées aux étapes précédentes : l'export s'arrête si elles diffèrent
EXPECTED_ROWS = {"FactSales": 109880, "DimCustomer": 93104, "DimProduct": 32081,
                 "DimGeography": 4270, "DimDate": 608}
EXPECTED_REVENUE = 13181027.13
EXPECTED_NEW_CUSTOMERS = 93104


def write_table(wb, name, df):
    """Écrit un DataFrame dans une feuille et le déclare comme table Excel."""
    ws = wb.create_sheet(name)
    ws.append(list(df.columns))
    for row in df.itertuples(index=False, name=None):
        ws.append([None if pd.isna(v) else v for v in row])

    # Les colonnes de dates reçoivent un vrai format date Excel
    for idx, col in enumerate(df.columns, start=1):
        if len(df) and isinstance(df[col].iloc[0], dt.date):
            for (cell,) in ws.iter_rows(min_row=2, min_col=idx, max_col=idx):
                cell.number_format = "yyyy-mm-dd"

    ref = f"A1:{get_column_letter(df.shape[1])}{len(df) + 1}"
    table = Table(displayName=name, ref=ref)
    table.tableStyleInfo = TableStyleInfo(name="TableStyleLight1", showRowStripes=False)
    ws.add_table(table)


def main():
    engine = get_engine()
    frames = {}
    for name, query in QUERIES.items():
        df = pd.read_sql(query, engine)
        assert len(df) == EXPECTED_ROWS[name], (
            f"{name} : {len(df)} lignes au lieu de {EXPECTED_ROWS[name]}")
        frames[name] = df
        print(f"{name:13s} {len(df):>8,} lignes x {df.shape[1]} colonnes")

    fact = frames["FactSales"]
    revenue = round(float(fact["revenue"].sum()), 2)
    new_customers = fact.loc[fact["is_first_purchase_day"] == 1, "customer_key"].nunique()
    assert abs(revenue - EXPECTED_REVENUE) < 0.01, revenue
    assert new_customers == EXPECTED_NEW_CUSTOMERS, new_customers
    print(f"Contrôles OK : CA = {revenue:,.2f} | nouveaux clients = {new_customers:,}")

    wb = Workbook()
    wb.remove(wb.active)
    for name, df in frames.items():
        write_table(wb, name, df)
    OUT_PATH.parent.mkdir(parents=True, exist_ok=True)
    wb.save(OUT_PATH)
    print(f"Classeur : {OUT_PATH} ({OUT_PATH.stat().st_size / 1e6:.1f} Mo)")


if __name__ == "__main__":
    main()
